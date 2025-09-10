import os
from pathlib import Path
from datetime import datetime, date
import csv
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText

from docx import Document
from openpyxl import load_workbook
from PyPDF2 import PdfReader

# Funktion für Autorenmetadata
def get_author(file_path):
    try:
        if file_path.suffix.lower() == ".docx":
            doc = Document(file_path)
            return doc.core_properties.author or "Unknown"
        elif file_path.suffix.lower() == ".xlsx":
            wb = load_workbook(file_path, read_only=True)
            return wb.properties.creator or "Unknown"
        elif file_path.suffix.lower() == ".pdf":
            pdf = PdfReader(str(file_path))
            return pdf.metadata.author or "Unknown"
        else:
            return "N/A"
    except:
        return "Error"

# Recursive file scan mit Filtern
def scan_files(folder_path, file_types):
    files = []
    for root, _, filenames in os.walk(folder_path):
        for name in filenames:
            file_path = Path(root) / name
            if not file_types or file_path.suffix.lower() in file_types:
                files.append(file_path)
    return files

# Funktion: Statusmeldung ins Log schreiben
def log_message(message):
    log_box.config(state="normal")
    log_box.insert(tk.END, message + "\n")
    log_box.see(tk.END)  # Automatisch nach unten scrollen
    log_box.config(state="disabled")
    root.update_idletasks()

# Metadaten extrahieren
def extract_metadata(folder_path, file_types, csv_save_path):
    files = scan_files(folder_path, file_types)
    total_files = len(files)
    if total_files == 0:
        messagebox.showinfo("Info", "Keine Dateien mit dem gewählten Filter gefunden.")
        return

    file_metadata = []

    progress_bar["maximum"] = total_files
    log_message(f"Starte Zusammenstellung in {folder_path} ({total_files} Dateien gefunden)")

    for idx, file in enumerate(files, 1):
        try:
            info = file.stat()
            file_name = file.name
            author = get_author(file)
            creation_date = datetime.fromtimestamp(info.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
            modification_date = datetime.fromtimestamp(info.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            access_date = datetime.fromtimestamp(info.st_atime).strftime("%Y-%m-%d %H:%M:%S")
            size_mb = round(info.st_size / (1024 * 1024), 2)

            file_metadata.append([file_name, author, size_mb, creation_date, modification_date, access_date])

            progress_bar["value"] = idx
            log_message(f"[{idx}/{total_files}] Fertig: {file_name}")

        except Exception as e:
            log_message(f"Fehler bei Datei {file}: {e}")

# Als CSV speichern
    with open(csv_save_path, mode='w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f, delimiter="|")
        writer.writerow(["Dateiname", "Autor", "Größe in MB", "Erstelldatum", "Änderungsdatum", "Letzter Zugriff"])
        writer.writerows(file_metadata)

    log_message(f"✅ Fertig! Datei gespeichert unter: {csv_save_path}")
    messagebox.showinfo("Fertig!", f"Liste wurde gespeichert in:\n{csv_save_path}")

# UI Funktionen
def browse_folder():
    folder_selected = filedialog.askdirectory()
    if folder_selected:
        folder_path_var.set(folder_selected)
        # Standard CSV-Pfad automatisch im gewählten Ordner setzen
        default_name = f"Übergabeprotokoll_{date.today().isoformat()}.csv"
        csv_file_default = os.path.abspath(os.path.join(folder_selected, default_name))
        csv_path_var.set(csv_file_default)

def browse_csv_file():
    file_selected = filedialog.asksaveasfilename(defaultextension=".csv",
                                                 filetypes=[("CSV files", "*.csv")])
    if file_selected:
        csv_path_var.set(file_selected)

def run_script():
    folder = folder_path_var.get()
    csv_file = csv_path_var.get()
    filters = [ft.strip().lower() for ft in file_types_var.get().split(",") if ft.strip()]
    filters = [f if f.startswith(".") else f".{f}" for f in filters]

    if not folder:
        messagebox.showwarning("Warnung", "Bitte wählen Sie einen Ordner aus.")
        return

    # Standard CSV-Datei im gleichen Ordner wie der Scan-Ordner
    if not csv_file:
        default_name = f"Übergabeprotokoll_{date.today().isoformat()}.csv"
        csv_file = os.path.abspath(os.path.join(folder, default_name))
        csv_path_var.set(csv_file)  # UI-Feld mit vollständigem Pfad aktualisieren
        log_message(f"Datei wird automatisch gespeichert unter: {csv_file}")
    # Log und Fortschritt zurücksetzen
    progress_bar["value"] = 0
    log_box.config(state="normal")
    log_box.delete("1.0", tk.END)
    log_box.config(state="disabled")

    # Metadaten extrahieren
    extract_metadata(folder, filters, csv_file)

# Tkinter UI
root = tk.Tk()
root.title("Digitaler Übernahmeassistent des AdsD")
root.geometry("600x420")
root.resizable(False, False)
#Farbe klappt im moment noch nicht, muss noch angepasst werden ?
root.configure(bg="#B22222")  # Farbe (FireBrick)

# Styles
style = ttk.Style()
style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=5)
style.configure("TLabel", font=("Segoe UI", 10))
style.configure("TEntry", font=("Segoe UI", 10))

# Variables
folder_path_var = tk.StringVar()
csv_path_var = tk.StringVar()
file_types_var = tk.StringVar()

# Frames
frame_top = ttk.Frame(root, padding=(10,10))
frame_top.pack(fill="x")
frame_top.configure(style="Red.TFrame")

frame_middle = ttk.Frame(root, padding=(10,5))
frame_middle.pack(fill="x")
frame_middle.configure(style="Red.TFrame")

frame_bottom = ttk.Frame(root, padding=(10,5))
frame_bottom.pack(fill="x")
frame_bottom.configure(style="Red.TFrame")

frame_log = ttk.Frame(root, padding=(10,5))
frame_log.pack(fill="both", expand=True)
frame_log.configure(style="Red.TFrame")

# Spaltenbreiten festlegen, damit Buttons immer gleich ausgerichtet sind
frame_top.columnconfigure(0, minsize=500)
frame_middle.columnconfigure(0, minsize=500)

# Top Frame
frame_folder = ttk.Frame(frame_top)
frame_folder.grid(row=1, column=0, columnspan=2, sticky="w")
# Eingabeordner
ttk.Label(frame_top, text="Übergabeordner oder -Laufwerk:").grid(row=0, column=0, sticky="w")
tk.Entry(frame_folder, textvariable=folder_path_var, width=60).pack(side="left")
ttk.Button(frame_folder, text="Auswählen", command=browse_folder).pack(side="left", padx=5)
# Filter
ttk.Label(frame_top, text="Optional: Dateitypen (z.B. .docx; mit Komma separieren; leerlassen um alle auszuwählen):").grid(row=2, column=0, sticky="w", pady=(10,0))
ttk.Entry(frame_top, textvariable=file_types_var, width=60).grid(row=3, column=0, sticky="w")
# Mittelframe
frame_csv = ttk.Frame(frame_middle)
frame_csv.grid(row=1, column=0, columnspan=2, sticky="w")
ttk.Label(frame_middle, text="Ergebnis speichern unter:").grid(row=0, column=0, sticky="w")
ttk.Entry(frame_csv, textvariable=csv_path_var, width=60).pack(side="left")
ttk.Button(frame_csv, text="Ausgabeziel ändern", command=browse_csv_file).pack(side="left", padx=5)

# --- Bottom Frame: Fortschritt + Start Button ---
progress_bar = ttk.Progressbar(frame_bottom, length=500, mode="determinate")
progress_bar.pack(pady=(5,10))
ttk.Button(frame_bottom, text="Übernahmeprotokoll erstellen", command=run_script).pack()

# --- Log Frame ---
log_box = ScrolledText(frame_log, width=40, height=2, state="disabled", font=("Segoe UI", 10))
log_box.pack(fill="both", expand=True)

root.mainloop()
